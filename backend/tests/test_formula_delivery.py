"""公式本身就是验证。

「用 VLOOKUP 把客户名称匹配过来」要的不是同一批值换条路算出来，而是一份**会自证**的
工作簿：点开填充列的单元格，看到 ``=VLOOKUP($B2,客户档案!$A:$B,2,FALSE)``，和用户自己
写的一模一样。没匹配上的行显示 #N/A —— 那正是 Excel 里该有的样子，一眼就能看见。

所以不需要「填充核对」sheet，不需要「问题说明」，不需要在旁边加一列匹配状态。多加的
每一张表都在降低交付物的可读性。

公式引用的源表必须在同一个工作簿里，否则单元格是坏的。那张表不是「多加的 sheet」，
它是这个请求本身的一部分。
"""

from __future__ import annotations

import warnings
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from data_agent.services import answer_input_paths
from data_agent.tools.excel_formula import lookup_formula


@pytest.fixture(autouse=True)
def deterministic(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("DATA_AGENT_LLM_ENABLED", "0")
    warnings.filterwarnings("ignore")


def _inputs(tmp_path: Path) -> Path:
    src = tmp_path / "in"
    src.mkdir(parents=True, exist_ok=True)
    pd.DataFrame(
        {"订单号": ["O1", "O2", "O3"], "客户编号": ["C1", "C2", "C9"], "金额": [100, 200, 300]}
    ).to_excel(src / "订单明细.xlsx", index=False)
    pd.DataFrame(
        {"客户编号": ["C1", "C2"], "客户名称": ["张三", "李四"], "所属大区": ["华北", "华东"]}
    ).to_excel(src / "客户档案.xlsx", index=False)
    return src


def _deliver(tmp_path: Path, goal: str):
    answer_input_paths([_inputs(tmp_path)], goal=goal, output_dir=tmp_path / "out")
    return load_workbook(tmp_path / "out" / "final_result.xlsx")


# --- 公式构造 -----------------------------------------------------------------


def test_vlookup_points_at_the_right_column() -> None:
    formula = lookup_formula(
        result_columns=["订单号", "客户编号", "客户名称"],
        left_key="客户编号",
        source_sheet="客户档案",
        source_columns=["客户编号", "客户名称", "所属大区"],
        right_key="客户编号",
        value_field="所属大区",
        row=2,
    )
    assert formula == "=VLOOKUP($B2,客户档案!$A:$C,3,FALSE)"


def test_a_value_left_of_the_key_uses_index_match() -> None:
    """VLOOKUP 只能向右找。Excel 用户遇到这种表会写 INDEX/MATCH。"""

    formula = lookup_formula(
        result_columns=["订单号", "客户编号"],
        left_key="客户编号",
        source_sheet="客户档案",
        source_columns=["客户名称", "客户编号"],
        right_key="客户编号",
        value_field="客户名称",
        row=2,
    )
    assert formula == "=INDEX(客户档案!$A:$A,MATCH($B2,客户档案!$B:$B,0))"


def test_a_sheet_name_with_spaces_is_quoted() -> None:
    formula = lookup_formula(
        result_columns=["客户编号"],
        left_key="客户编号",
        source_sheet="客户 档案",
        source_columns=["客户编号", "客户名称"],
        right_key="客户编号",
        value_field="客户名称",
        row=2,
    )
    assert "'客户 档案'!" in formula


# --- 交付物 -------------------------------------------------------------------


def test_the_filled_column_holds_real_formulas(tmp_path: Path) -> None:
    workbook = _deliver(tmp_path, "用VLOOKUP把客户名称匹配到订单明细")
    sheet = workbook["处理结果"]
    header = [cell.value for cell in sheet[1]]
    column = header.index("客户名称") + 1

    for row in (2, 3, 4):
        cell = sheet.cell(row=row, column=column)
        assert cell.data_type == "f", "写成了文本，Excel 打开不会计算"
        assert cell.value.startswith("=VLOOKUP(")


def test_the_source_table_ships_with_the_formula(tmp_path: Path) -> None:
    """公式指向的表不在工作簿里，单元格就是坏的。"""

    workbook = _deliver(tmp_path, "用VLOOKUP把客户名称匹配到订单明细")
    assert "客户档案" in workbook.sheetnames


def test_no_extra_sheets_beyond_what_the_request_needs(tmp_path: Path) -> None:
    """交付物只有结果表和公式的数据源 —— 没有核对表、没有问题说明。"""

    workbook = _deliver(tmp_path, "用VLOOKUP把客户名称匹配到订单明细")
    assert set(workbook.sheetnames) == {"处理结果", "客户档案"}


def test_an_ordinary_lookup_still_delivers_values(tmp_path: Path) -> None:
    """没要求公式就还是值。公式是用户点名要的东西，不是新的默认。"""

    workbook = _deliver(tmp_path, "把客户名称关联到订单明细")
    sheet = workbook["处理结果"]
    header = [cell.value for cell in sheet[1]]
    column = header.index("客户名称") + 1

    assert sheet.cell(row=2, column=column).data_type != "f"
    assert "客户档案" not in workbook.sheetnames
