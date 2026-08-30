import pandas as pd
from openpyxl import load_workbook

from data_agent.tools.exporter import export_workbook


def test_export_workbook_handles_long_name_collisions(tmp_path) -> None:
    # Two names sharing a 31-char prefix must not overwrite each other.
    name_a = "very_long_sheet_name_that_exceeds_limit_alpha"
    name_b = "very_long_sheet_name_that_exceeds_limit_beta"
    output = tmp_path / "collide.xlsx"

    export_workbook(
        output,
        {
            name_a: pd.DataFrame({"a": [1]}),
            name_b: pd.DataFrame({"b": [2]}),
        },
    )

    sheet_names = pd.ExcelFile(output).sheet_names
    assert len(sheet_names) == 2
    assert len(set(sheet_names)) == 2
    for name in sheet_names:
        assert len(name) <= 31


def test_export_workbook_preserves_short_names(tmp_path) -> None:
    output = tmp_path / "simple.xlsx"
    export_workbook(
        output,
        {"处理结果": pd.DataFrame({"x": [1]}), "问题说明": pd.DataFrame({"y": [2]})},
    )

    assert pd.ExcelFile(output).sheet_names == ["处理结果", "问题说明"]


def test_export_workbook_applies_business_readability_formatting(tmp_path) -> None:
    output = tmp_path / "formatted.xlsx"
    export_workbook(
        output,
        {
            "处理结果": pd.DataFrame(
                {
                    "处理状态": ["通过", "需处理"],
                    "客户名称": ["北京示例客户", "上海示例客户"],
                }
            ),
            "问题说明": pd.DataFrame({"问题类型": ["金额为空"]}),
        },
    )

    workbook = load_workbook(output)
    result_sheet = workbook["处理结果"]
    problem_sheet = workbook["问题说明"]

    assert result_sheet.freeze_panes == "A2"
    assert result_sheet.auto_filter.ref == "A1:B3"
    assert result_sheet.sheet_view.showGridLines is False
    assert result_sheet["A1"].font.bold is True
    assert result_sheet["A1"].font.color.rgb == "00FFFFFF"
    assert result_sheet.column_dimensions["B"].width >= 12
    assert result_sheet.sheet_properties.tabColor.rgb == "000F766E"
    assert problem_sheet.sheet_properties.tabColor.rgb == "00D97706"
    assert len(result_sheet.conditional_formatting) == 2


def test_export_workbook_keeps_untrusted_formula_like_values_as_text(tmp_path) -> None:
    output = tmp_path / "untrusted-formulas.xlsx"
    values = ["=1+1", "+cmd", "-2+3", "@SUM(A1:A2)", "＝1+1"]

    export_workbook(output, {"处理结果": pd.DataFrame({"用户输入": values})})

    sheet = load_workbook(output, data_only=False)["处理结果"]
    for row, expected in enumerate(values, start=2):
        cell = sheet.cell(row=row, column=1)
        assert cell.value == expected
        assert cell.data_type == "s"


def test_only_explicitly_trusted_formula_columns_remain_executable(tmp_path) -> None:
    output = tmp_path / "trusted-formulas.xlsx"
    frame = pd.DataFrame(
        {
            "用户输入": ["=WEBSERVICE(\"https://example.invalid\")"],
            "客户名称": ["=VLOOKUP($A2,客户档案!$A:$B,2,FALSE)"],
        }
    )

    export_workbook(
        output,
        {"处理结果": frame},
        trusted_formula_columns={"处理结果": {"客户名称"}},
    )

    sheet = load_workbook(output, data_only=False)["处理结果"]
    assert sheet["A2"].data_type == "s"
    assert sheet["A2"].value == '=WEBSERVICE("https://example.invalid")'
    assert sheet["B2"].data_type == "f"
    assert sheet["B2"].value.startswith("=VLOOKUP(")
