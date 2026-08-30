from __future__ import annotations

from collections.abc import Collection, Mapping
from pathlib import Path
from unicodedata import east_asian_width

import pandas as pd
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

EXCEL_SHEET_NAME_LIMIT = 31
_HEADER_FILL = PatternFill("solid", fgColor="17324D")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_EVEN_ROW_FILL = PatternFill("solid", fgColor="F5F8F7")
_PASS_FILL = PatternFill("solid", fgColor="DCFCE7")
_PASS_FONT = Font(color="15803D", bold=True)
_REVIEW_FILL = PatternFill("solid", fgColor="FEF3C7")
_REVIEW_FONT = Font(color="B45309", bold=True)
_FAIL_FILL = PatternFill("solid", fgColor="FEE2E2")
_FAIL_FONT = Font(color="B91C1C", bold=True)
_FORMULA_LIKE_PREFIXES = ("=", "+", "-", "@", "\t", "\r", "\n", "＝", "＋", "－", "＠")


def _safe_sheet_name(raw_name: str, used_names: set[str]) -> str:
    """Truncate to Excel's 31-char limit and de-duplicate collisions.

    Two long names sharing a 31-char prefix would otherwise silently overwrite each
    other. When a truncated name is already taken, append a numeric suffix that keeps
    the result within the limit.
    """

    base = str(raw_name)[:EXCEL_SHEET_NAME_LIMIT] or "sheet"
    if base not in used_names:
        used_names.add(base)
        return base
    counter = 2
    while True:
        suffix = f"_{counter}"
        candidate = base[: EXCEL_SHEET_NAME_LIMIT - len(suffix)] + suffix
        if candidate not in used_names:
            used_names.add(candidate)
            return candidate
        counter += 1


def export_workbook(
    output_file: str | Path,
    sheets: dict[str, pd.DataFrame],
    *,
    trusted_formula_columns: Mapping[str, Collection[str]] | None = None,
) -> Path:
    output_path = Path(output_file)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    used_names: set[str] = set()
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        for sheet_name, sheet_df in sheets.items():
            safe_name = _safe_sheet_name(sheet_name, used_names)
            sheet_df.to_excel(writer, sheet_name=safe_name, index=False)
            _neutralize_untrusted_formulas(
                writer.sheets[safe_name],
                trusted_columns=set((trusted_formula_columns or {}).get(sheet_name, ())),
            )
            _format_worksheet(writer.sheets[safe_name], sheet_df, safe_name)

    return output_path


def _format_worksheet(worksheet, frame: pd.DataFrame, sheet_name: str) -> None:
    """Apply compact business-friendly formatting without changing workbook data."""

    worksheet.freeze_panes = "A2"
    worksheet.sheet_view.showGridLines = False
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.row_dimensions[1].height = 27
    worksheet.sheet_properties.tabColor = (
        "0F766E" if sheet_name == "处理结果" else "D97706"
        if sheet_name == "问题说明"
        else "54718C"
    )

    for cell in worksheet[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    sample = frame.head(200)
    for column_index, column_name in enumerate(frame.columns, start=1):
        values = [column_name, *sample.iloc[:, column_index - 1].tolist()]
        width = min(42, max(12, max(_display_width(value) for value in values) + 2))
        worksheet.column_dimensions[get_column_letter(column_index)].width = width

    if worksheet.max_row < 2 or worksheet.max_column < 1:
        return

    data_range = (
        f"A2:{get_column_letter(worksheet.max_column)}{worksheet.max_row}"
    )
    worksheet.conditional_formatting.add(
        data_range,
        FormulaRule(formula=["MOD(ROW(),2)=0"], fill=_EVEN_ROW_FILL),
    )

    status_column = _column_index(frame, "处理状态")
    if status_column is None:
        return
    letter = get_column_letter(status_column)
    status_range = f"{letter}2:{letter}{worksheet.max_row}"
    for value, fill, font in (
        ("通过", _PASS_FILL, _PASS_FONT),
        ("需处理", _REVIEW_FILL, _REVIEW_FONT),
        ("不通过", _FAIL_FILL, _FAIL_FONT),
    ):
        worksheet.conditional_formatting.add(
            status_range,
            FormulaRule(
                formula=[f'${letter}2="{value}"'],
                fill=fill,
                font=font,
            ),
        )


def _neutralize_untrusted_formulas(worksheet, *, trusted_columns: set[str]) -> None:
    """Keep uploaded formula-like text inert while preserving system formulas.

    The value itself is not escaped or normalized. Marking the XLSX cell as a string
    preserves what the user uploaded and prevents spreadsheet applications from
    evaluating it. Only columns explicitly identified by the deterministic runner may
    contain formulas; callers that merely export data have no trusted columns.
    """

    column_names = {
        index: str(worksheet.cell(row=1, column=index).value or "")
        for index in range(1, worksheet.max_column + 1)
    }
    for row in worksheet.iter_rows():
        for cell in row:
            value = cell.value
            if not isinstance(value, str) or not value.startswith(_FORMULA_LIKE_PREFIXES):
                continue
            is_trusted_formula = (
                cell.row > 1
                and column_names.get(cell.column) in trusted_columns
                and value.startswith("=")
            )
            if not is_trusted_formula:
                cell.data_type = "s"


def _display_width(value: object) -> int:
    if value is None or (pd.api.types.is_scalar(value) and pd.isna(value)):
        return 0
    text = str(value)
    return max(
        (
            sum(2 if east_asian_width(char) in {"W", "F"} else 1 for char in line)
            for line in text.splitlines()
        ),
        default=0,
    )


def _column_index(frame: pd.DataFrame, name: str) -> int | None:
    for index, column in enumerate(frame.columns, start=1):
        if str(column) == name:
            return index
    return None


def export_profile_report(
    output_file: str | Path,
    profile_sheets: dict[str, pd.DataFrame],
) -> Path:
    output_path = Path(output_file)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    used_names: set[str] = set()
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        for sheet_name, sheet_df in profile_sheets.items():
            safe_name = _safe_sheet_name(sheet_name, used_names)
            sheet_df.to_excel(writer, sheet_name=safe_name, index=False)
            _neutralize_untrusted_formulas(writer.sheets[safe_name], trusted_columns=set())

    return output_path
