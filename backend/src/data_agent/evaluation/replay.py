"""Offline replay of versioned golden tasks through the production pipeline."""

from __future__ import annotations

import hashlib
import json
import os
import re
from collections.abc import Iterator
from contextlib import contextmanager
from datetime import datetime, timezone
from pathlib import Path
from time import perf_counter
from typing import Any, Optional

import pandas as pd
from openpyxl import load_workbook

from data_agent.agent.llm_client import load_llm_config
from data_agent.agent.prompts import GOAL_PROMPT_VERSION, PLANNER_PROMPT_VERSION
from data_agent.capabilities import PLAN_COMPILER_VERSION, get_capability_registry
from data_agent.evaluation.golden import GoldenTask, load_golden_tasks
from data_agent.evaluation.metrics import aggregate_replay_metrics, coverage
from data_agent.observability import capture_llm_usage
from data_agent.services import (
    build_plan_snapshot,
    execute_planned_job,
    hydrate_plan_snapshot,
    plan_input_paths,
)


def run_golden_suite(
    output_root: str | Path,
    *,
    tasks: Optional[list[GoldenTask]] = None,
    run_id: Optional[str] = None,
    planning_mode: str = "deterministic",
    cache_mode: str = "cold",
) -> dict[str, Any]:
    if planning_mode not in {"deterministic", "configured"}:
        raise ValueError("planning_mode must be deterministic or configured")
    if cache_mode not in {"off", "cold", "warm"}:
        raise ValueError("cache_mode must be off, cold or warm")
    selected = tasks or load_golden_tasks()
    requested_id = run_id or datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    identifier, run_dir = _allocate_run_directory(output_root, requested_id)

    with _evaluation_environment(
        run_dir,
        planning_mode=planning_mode,
        cache_mode=cache_mode,
    ):
        results = [
            (
                _skipped_task(task, "requires configured LLM")
                if planning_mode == "deterministic" and task.requires_llm
                else _evaluate_task(
                    task,
                    run_dir / task.task_id,
                    warm_cache=cache_mode == "warm",
                    expected_understanding_source=(
                        "llm_cached"
                        if planning_mode == "configured" and cache_mode == "warm"
                        else "llm"
                        if planning_mode == "configured"
                        else None
                    ),
                )
            )
            for task in selected
        ]
    report = {
        "run_id": identifier,
        "created_at": datetime.now(timezone.utc).isoformat(),
        "versions": _version_metadata(),
        "planning_mode": planning_mode,
        "cache_mode": cache_mode,
        "metrics": aggregate_replay_metrics(results),
        "results": results,
    }
    report_path = run_dir / "replay_report.json"
    report_path.write_text(
        json.dumps(report, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    report["report_path"] = str(report_path)
    return report


def rescore_golden_report(
    report_path: str | Path,
    output_root: str | Path,
    *,
    run_id: str | None = None,
) -> dict[str, Any]:
    """Re-evaluate persisted workbooks after expectation-only suite changes.

    This never plans or executes a task and therefore never calls the configured
    model. Runtime versions must match because artifact-only rescoring is valid for
    benchmark assertion changes, not for planner or prompt changes.
    """

    source_path = Path(report_path).resolve()
    source = json.loads(source_path.read_text(encoding="utf-8"))
    current_versions = _version_metadata()
    for key in (
        "goal_prompt",
        "planner_prompt",
        "plan_compiler",
        "capability_registry",
        "source_tree",
        "model",
    ):
        source_version = source.get("versions", {}).get(key)
        if source_version != current_versions[key]:
            raise ValueError(
                f"cannot rescore artifacts across {key} versions: "
                f"{source_version!r} != {current_versions[key]!r}"
            )

    tasks = {task.task_id: task for task in load_golden_tasks()}
    rescored_results: list[dict[str, Any]] = []
    for persisted in source.get("results", []):
        if persisted.get("skipped"):
            rescored_results.append(dict(persisted))
            continue
        task_id = str(persisted.get("task_id") or "")
        task = tasks.get(task_id)
        if task is None:
            raise ValueError(f"saved report references an unknown golden task: {task_id}")
        workbook_path = source_path.parent / task_id / "job" / "output" / "final_result.xlsx"
        if not workbook_path.is_file():
            raise FileNotFoundError(f"saved workbook is missing: {workbook_path}")
        rescored_results.append(_rescore_persisted_task(task, persisted, workbook_path))

    requested_id = run_id or f"{source.get('run_id', 'golden')}-rescored"
    identifier, run_dir = _allocate_run_directory(output_root, requested_id)
    report = {
        "run_id": identifier,
        "created_at": datetime.now(timezone.utc).isoformat(),
        "versions": current_versions,
        "planning_mode": source.get("planning_mode"),
        "cache_mode": source.get("cache_mode"),
        "rescored_from": str(source_path),
        "rescore": {"mode": "artifact_only", "llm_calls": 0},
        "metrics": aggregate_replay_metrics(rescored_results),
        "results": rescored_results,
    }
    destination = run_dir / "replay_report.json"
    destination.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    report["report_path"] = str(destination)
    return report


def merge_golden_reports(
    base_report_path: str | Path,
    replacement_report_paths: list[str | Path],
    output_root: str | Path,
    *,
    run_id: str | None = None,
) -> dict[str, Any]:
    """Create one release report from a full run plus targeted reruns.

    Only task ids present in the base report may be replaced. All reports must have
    identical version, planning-mode and cache-mode metadata, so a targeted rerun
    cannot silently mix a different model or evaluation contract into the gate.
    """

    base_path = Path(base_report_path).resolve()
    base = json.loads(base_path.read_text(encoding="utf-8"))
    sources = [(base_path, base)]
    for value in replacement_report_paths:
        path = Path(value).resolve()
        sources.append((path, json.loads(path.read_text(encoding="utf-8"))))

    for path, candidate in sources[1:]:
        if candidate.get("versions") != base.get("versions"):
            raise ValueError(f"replacement report has different versions: {path}")
        if candidate.get("planning_mode") != base.get("planning_mode"):
            raise ValueError(f"replacement report has different planning mode: {path}")
        if candidate.get("cache_mode") != base.get("cache_mode"):
            raise ValueError(f"replacement report has different cache mode: {path}")

    base_results = list(base.get("results") or [])
    base_ids = {str(item.get("task_id") or "") for item in base_results}
    replacements: dict[str, tuple[dict[str, Any], Path]] = {}
    for path, candidate in sources[1:]:
        for result in candidate.get("results") or []:
            task_id = str(result.get("task_id") or "")
            if task_id not in base_ids:
                raise ValueError(f"replacement task is absent from base report: {task_id}")
            replacements[task_id] = (dict(result), path.parent / task_id)

    merged_results: list[dict[str, Any]] = []
    for result in base_results:
        task_id = str(result.get("task_id") or "")
        replacement = replacements.get(task_id)
        if replacement is None:
            merged = dict(result)
            artifact_source = base_path.parent / task_id
        else:
            merged, artifact_source = replacement
        merged["artifact_source"] = str(artifact_source)
        merged_results.append(merged)

    requested_id = run_id or f"{base.get('run_id', 'golden')}-merged"
    identifier, run_dir = _allocate_run_directory(output_root, requested_id)
    report = {
        "run_id": identifier,
        "created_at": datetime.now(timezone.utc).isoformat(),
        "versions": base.get("versions"),
        "planning_mode": base.get("planning_mode"),
        "cache_mode": base.get("cache_mode"),
        "merged_from": [str(path) for path, _report in sources],
        "merge": {
            "mode": "targeted_replacement",
            "llm_calls": 0,
            "replaced_task_ids": sorted(replacements),
        },
        "metrics": aggregate_replay_metrics(merged_results),
        "results": merged_results,
    }
    destination = run_dir / "replay_report.json"
    destination.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    report["report_path"] = str(destination)
    return report


def _rescore_persisted_task(
    task: GoldenTask,
    persisted: dict[str, Any],
    workbook_path: Path,
) -> dict[str, Any]:
    """Apply current Golden expectations to one immutable saved workbook."""

    expected = task.expectations
    actual = dict(persisted.get("actual") or {})
    actual_actions = set(actual.get("task_actions") or [])
    actual_capabilities = set(actual.get("capabilities") or [])
    actual_sheets = pd.ExcelFile(workbook_path).sheet_names
    sheet_expectations = expected.sheet_expectations or {
        expected.sheets[0]: {
            "required_columns": expected.required_columns,
            "forbidden_columns": expected.forbidden_columns,
            "row_assertions": expected.row_assertions,
            "forbidden_rows": expected.forbidden_rows,
            "min_rows": 0,
        }
    }
    per_sheet_checks: dict[str, dict[str, bool]] = {}
    for sheet_name, expectation in sheet_expectations.items():
        if sheet_name not in actual_sheets:
            per_sheet_checks[sheet_name] = {
                "exact_columns": False,
                "required_columns": False,
                "forbidden_columns": False,
                "row_assertions": False,
                "forbidden_rows": False,
                "min_rows": False,
                "exact_rows": False,
            }
            continue
        frame = pd.read_excel(workbook_path, sheet_name=sheet_name)
        per_sheet_checks[sheet_name] = _sheet_checks(frame, expectation)

    previous_checks = dict(persisted.get("checks") or {})
    checks = {
        "task_actions": set(expected.task_actions) <= actual_actions,
        "forbidden_task_actions": set(expected.forbidden_task_actions).isdisjoint(
            actual_actions
        ),
        "capabilities": set(expected.capabilities) <= actual_capabilities,
        "forbidden_capabilities": set(expected.forbidden_capabilities).isdisjoint(
            actual_capabilities
        ),
        "sheets": actual_sheets == expected.sheets,
        "charts": list(actual.get("charts") or []) == expected.charts,
        "report": bool(actual.get("report_enabled")) is expected.report_enabled,
        "clarification_slots": int(actual.get("clarification_slots", 0))
        <= expected.max_clarification_slots,
        # These are runtime observations, not benchmark expectations. Preserve them
        # because rescoring deliberately does not execute the pipeline again.
        "execution_status": bool(previous_checks.get("execution_status")),
        "counts": all(
            int((actual.get("counts") or {}).get(key, -1)) == expected_value
            for key, expected_value in expected.counts.items()
        ),
        "required_columns": all(
            item["required_columns"] for item in per_sheet_checks.values()
        ),
        "forbidden_columns": all(
            item["forbidden_columns"] for item in per_sheet_checks.values()
        ),
        "row_assertions": all(
            item["row_assertions"] for item in per_sheet_checks.values()
        ),
        "forbidden_rows": all(
            item["forbidden_rows"] for item in per_sheet_checks.values()
        ),
        "sheet_row_counts": all(item["min_rows"] for item in per_sheet_checks.values()),
        "exact_sheet_rows": all(
            item["exact_rows"] for item in per_sheet_checks.values()
        ),
        "exact_columns": all(
            item["exact_columns"] for item in per_sheet_checks.values()
        ),
        "formula_safety": expected.allow_formulas
        or not _workbook_contains_formula(workbook_path),
        "plan_snapshot": bool(previous_checks.get("plan_snapshot")),
    }
    if "understanding_source" in previous_checks:
        checks["understanding_source"] = bool(previous_checks["understanding_source"])

    output_checks = (checks["sheets"], checks["charts"], checks["report"])
    result_checks = (
        checks["counts"],
        checks["required_columns"],
        checks["forbidden_columns"],
        checks["row_assertions"],
        checks["forbidden_rows"],
        checks["sheet_row_counts"],
        checks["exact_sheet_rows"],
        checks["exact_columns"],
        checks["formula_safety"],
    )
    rescored = dict(persisted)
    rescored["passed"] = all(checks.values())
    rescored["checks"] = checks
    rescored["metrics"] = {
        "intent_coverage": coverage(set(expected.task_actions), actual_actions),
        "capability_coverage": coverage(
            set(expected.capabilities), actual_capabilities
        ),
        "output_compliance": round(sum(output_checks) / len(output_checks), 4),
        "result_semantics": round(sum(result_checks) / len(result_checks), 4),
    }
    rescored["actual"] = {
        **actual,
        "sheets": actual_sheets,
        "sheet_checks": per_sheet_checks,
    }
    return rescored


def _allocate_run_directory(
    output_root: str | Path,
    requested_id: str,
) -> tuple[str, Path]:
    """Reserve a non-destructive run directory, suffixing repeated identifiers."""

    if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9._-]{0,127}", requested_id):
        raise ValueError(
            "run_id 只能包含字母、数字、点、下划线和连字符，且必须以字母或数字开头"
        )
    root = Path(output_root)
    root.mkdir(parents=True, exist_ok=True)
    attempt = 1
    while True:
        identifier = requested_id if attempt == 1 else f"{requested_id}-{attempt}"
        run_dir = root / identifier
        try:
            run_dir.mkdir()
        except FileExistsError:
            attempt += 1
            continue
        return identifier, run_dir


def run_golden_task(
    task: GoldenTask,
    task_dir: str | Path,
    *,
    warm_cache: bool = False,
) -> dict[str, Any]:
    started = perf_counter()
    work_path = Path(task_dir)
    input_dir = work_path / "input"
    input_dir.mkdir(parents=True, exist_ok=False)
    input_paths = []
    all_tables = {
        **task.tables,
        **{
            table_name: _generate_table(spec)
            for table_name, spec in task.generated_tables.items()
        },
    }
    for table_name, columns in all_tables.items():
        options = task.input_options.get(table_name)
        extension = "xlsx" if options and options.file_format == "xlsx" else "csv"
        path = input_dir / f"{table_name}.{extension}"
        frame = pd.DataFrame(columns)
        if extension == "xlsx":
            _write_excel_fixture(path, frame, options)
        else:
            frame.to_csv(path, index=False)
        input_paths.append(path)

    try:
        if warm_cache:
            plan_input_paths(
                input_paths,
                work_path / "warmup",
                config={"goal": task.goal, "include_diagnostics": True},
            )
        plan = plan_input_paths(
            input_paths,
            work_path / "job",
            config={"goal": task.goal, "include_diagnostics": True},
        )
        snapshot = build_plan_snapshot(plan, input_paths)
        hydrated = hydrate_plan_snapshot(snapshot, work_path / "job")
        response = execute_planned_job(hydrated, input_paths)

        task_spec = plan.goal_plan.get("task_spec") or {}
        actual_actions = set(task_spec.get("actions") or [])
        actual_capabilities = {
            step.capability_id for step in plan.execution_plan.steps
        }
        expected = task.expectations
        actual_sheets = pd.ExcelFile(response["files"]["excel_path"]).sheet_names
        workbook_path = response["files"]["excel_path"]
        sheet_expectations = expected.sheet_expectations or {
            expected.sheets[0]: {
                "required_columns": expected.required_columns,
                "forbidden_columns": expected.forbidden_columns,
                "row_assertions": expected.row_assertions,
                "forbidden_rows": expected.forbidden_rows,
                "min_rows": 0,
            }
        }
        frames = {
            sheet_name: pd.read_excel(workbook_path, sheet_name=sheet_name)
            for sheet_name in sheet_expectations
        }
        per_sheet_checks = {
            sheet_name: _sheet_checks(frames[sheet_name], expectation)
            for sheet_name, expectation in sheet_expectations.items()
        }
        actual_charts = [
            str(item.get("chart_id")) for item in response.get("charts", [])
        ]
        actual_report = any(
            key.startswith("report_") for key in response.get("files", {})
        )
        blocking_clarification_slots = sum(
            1
            for slot in task_spec.get("missing_slots") or []
            if isinstance(slot, dict) and slot.get("priority") == "high"
        )
        expected_actions = set(expected.task_actions)
        forbidden_actions = set(expected.forbidden_task_actions)
        expected_capabilities = set(expected.capabilities)
        forbidden_capabilities = set(expected.forbidden_capabilities)
        checks = {
            "task_actions": expected_actions <= actual_actions,
            "forbidden_task_actions": forbidden_actions.isdisjoint(actual_actions),
            "capabilities": expected_capabilities <= actual_capabilities,
            "forbidden_capabilities": forbidden_capabilities.isdisjoint(
                actual_capabilities
            ),
            "sheets": actual_sheets == expected.sheets,
            "charts": actual_charts == expected.charts,
            "report": actual_report is expected.report_enabled,
            "clarification_slots": blocking_clarification_slots
            <= expected.max_clarification_slots,
            "execution_status": response.get("status") == "success",
            "counts": all(
                int(response.get("counts", {}).get(key, -1)) == expected_value
                for key, expected_value in expected.counts.items()
            ),
            "required_columns": all(
                item["required_columns"] for item in per_sheet_checks.values()
            ),
            "forbidden_columns": all(
                item["forbidden_columns"] for item in per_sheet_checks.values()
            ),
            "row_assertions": all(
                item["row_assertions"] for item in per_sheet_checks.values()
            ),
            "forbidden_rows": all(
                item["forbidden_rows"] for item in per_sheet_checks.values()
            ),
            "sheet_row_counts": all(
                item["min_rows"] for item in per_sheet_checks.values()
            ),
            "exact_sheet_rows": all(
                item["exact_rows"] for item in per_sheet_checks.values()
            ),
            "exact_columns": all(
                item["exact_columns"] for item in per_sheet_checks.values()
            ),
            "formula_safety": expected.allow_formulas
            or not _workbook_contains_formula(workbook_path),
            "plan_snapshot": (
                snapshot.execution_plan == hydrated.execution_plan
                and snapshot.output_spec == hydrated.output_spec
                and snapshot.goal_plan == hydrated.goal_plan
            ),
        }
        output_checks = (
            checks["sheets"],
            checks["charts"],
            checks["report"],
        )
        result_checks = (
            checks["counts"],
            checks["required_columns"],
            checks["forbidden_columns"],
            checks["row_assertions"],
            checks["forbidden_rows"],
            checks["sheet_row_counts"],
            checks["exact_sheet_rows"],
            checks["exact_columns"],
            checks["formula_safety"],
        )
        return {
            "task_id": task.task_id,
            "scenario_id": task.scenario_id or task.task_id,
            "variant_index": task.variant_index,
            "tags": task.tags,
            "domain": task.domain,
            "passed": all(checks.values()),
            "checks": checks,
            "metrics": {
                "intent_coverage": coverage(expected_actions, actual_actions),
                "capability_coverage": coverage(
                    expected_capabilities,
                    actual_capabilities,
                ),
                "output_compliance": round(
                    sum(output_checks) / len(output_checks),
                    4,
                ),
                "result_semantics": round(
                    sum(result_checks) / len(result_checks),
                    4,
                ),
            },
            "actual": {
                "task_actions": sorted(actual_actions),
                "capabilities": [
                    step.capability_id for step in plan.execution_plan.steps
                ],
                "sheets": actual_sheets,
                "charts": actual_charts,
                "report_enabled": actual_report,
                "clarification_slots": blocking_clarification_slots,
                "non_blocking_clarification_slots": max(
                    0,
                    len(task_spec.get("missing_slots") or [])
                    - blocking_clarification_slots,
                ),
                "understanding_source": plan.goal_plan.get("understanding_source"),
                "llm_attempts": plan.llm_attempts.to_dict(orient="records"),
                "plan_hash": snapshot.plan_hash,
                "counts": response.get("counts", {}),
                "sheet_checks": per_sheet_checks,
            },
            "duration_ms": round((perf_counter() - started) * 1000, 2),
            "error": None,
        }
    except Exception as exc:  # noqa: BLE001 - replay records failures per task
        return {
            "task_id": task.task_id,
            "scenario_id": task.scenario_id or task.task_id,
            "variant_index": task.variant_index,
            "tags": task.tags,
            "domain": task.domain,
            "passed": False,
            "checks": {},
            "metrics": {
                "intent_coverage": 0.0,
                "capability_coverage": 0.0,
                "output_compliance": 0.0,
                "result_semantics": 0.0,
            },
            "actual": {},
            "duration_ms": round((perf_counter() - started) * 1000, 2),
            "error": f"{type(exc).__name__}: {exc}",
        }


def _evaluate_task(
    task: GoldenTask,
    task_dir: Path,
    *,
    warm_cache: bool,
    expected_understanding_source: str | None,
) -> dict[str, Any]:
    with capture_llm_usage() as usage:
        result = run_golden_task(task, task_dir, warm_cache=warm_cache)
    result["llm_usage"] = list(usage)
    if expected_understanding_source and result.get("checks"):
        matched = (
            result.get("actual", {}).get("understanding_source")
            == expected_understanding_source
        )
        result["checks"]["understanding_source"] = matched
        result["passed"] = bool(result.get("passed")) and matched
    return result


def _skipped_task(task: GoldenTask, reason: str) -> dict[str, Any]:
    return {
        "task_id": task.task_id,
        "scenario_id": task.scenario_id or task.task_id,
        "variant_index": task.variant_index,
        "tags": task.tags,
        "domain": task.domain,
        "passed": None,
        "skipped": True,
        "checks": {},
        "metrics": {},
        "actual": {},
        "duration_ms": 0.0,
        "error": None,
        "skip_reason": reason,
    }


def _generate_table(spec: Any) -> dict[str, list[Any]]:
    columns: dict[str, list[Any]] = {}
    for name, generator in spec.columns.items():
        if generator.kind == "sequence":
            columns[name] = [
                (
                    f"{generator.prefix}{generator.start + index * generator.step}"
                    if generator.prefix
                    else generator.start + index * generator.step
                )
                for index in range(spec.row_count)
            ]
        else:
            columns[name] = [
                generator.low if index < generator.low_count else generator.high
                for index in range(spec.row_count)
            ]
    return columns


def _write_excel_fixture(path: Path, frame: pd.DataFrame, options: Any) -> None:
    if options and options.header_rows:
        rows = [*options.header_rows, *frame.astype(object).values.tolist()]
        pd.DataFrame(rows).to_excel(path, index=False, header=False)
    else:
        frame.to_excel(path, index=False)
    if not options or (not options.merge_cells and not options.hidden_rows):
        return
    workbook = load_workbook(path)
    sheet = workbook.active
    for cell_range in options.merge_cells:
        sheet.merge_cells(cell_range)
    for row_index in options.hidden_rows:
        sheet.row_dimensions[row_index].hidden = True
    workbook.save(path)


def compare_replay_reports(
    current: dict[str, Any],
    baseline: dict[str, Any],
) -> dict[str, Any]:
    baseline_by_task = {
        item["task_id"]: item for item in baseline.get("results", [])
    }
    current_by_task = {
        item["task_id"]: item for item in current.get("results", [])
    }
    regressions = []
    for result in current.get("results", []):
        if result.get("skipped"):
            continue
        previous = baseline_by_task.get(result["task_id"])
        if previous and previous.get("passed") and not result.get("passed"):
            regressions.append(result["task_id"])
    missing_tasks = sorted(
        task_id
        for task_id, previous in baseline_by_task.items()
        if previous.get("passed") and task_id not in current_by_task
    )
    regressions.extend(missing_tasks)
    metric_names = (
        "pass_rate",
        "intent_coverage",
        "capability_coverage",
        "output_compliance",
        "result_semantics",
    )
    deltas = {
        name: round(
            float(current.get("metrics", {}).get(name, 0.0))
            - float(baseline.get("metrics", {}).get(name, 0.0)),
            4,
        )
        for name in metric_names
    }
    return {
        "regressions": regressions,
        "missing_tasks": missing_tasks,
        "metric_deltas": deltas,
        "passed": not regressions and all(value >= 0 for value in deltas.values()),
    }


def _sheet_checks(frame: pd.DataFrame, expectation: Any) -> dict[str, bool]:
    if hasattr(expectation, "model_dump"):
        expected = expectation.model_dump(mode="python")
    else:
        expected = dict(expectation)
    return {
        "exact_columns": (
            not expected.get("exact_columns")
            or list(frame.columns) == expected["exact_columns"]
        ),
        "required_columns": set(expected.get("required_columns", []))
        <= set(frame.columns),
        "forbidden_columns": set(expected.get("forbidden_columns", [])).isdisjoint(
            frame.columns
        ),
        "row_assertions": all(
            _contains_row(frame, assertion)
            for assertion in expected.get("row_assertions", [])
        ),
        "forbidden_rows": all(
            not _contains_row(frame, assertion)
            for assertion in expected.get("forbidden_rows", [])
        ),
        "min_rows": len(frame) >= int(expected.get("min_rows", 0)),
        "exact_rows": (
            expected.get("exact_rows") is None
            or len(frame) == int(expected["exact_rows"])
        ),
    }


def _contains_row(frame: pd.DataFrame, assertion: dict[str, Any]) -> bool:
    if not assertion or any(field not in frame.columns for field in assertion):
        return False
    mask = pd.Series(True, index=frame.index)
    for field, expected in assertion.items():
        values = frame[field]
        if expected is None:
            mask &= values.isna()
        else:
            mask &= values.astype(str).eq(str(expected))
    return bool(mask.any())


def _version_metadata() -> dict[str, str]:
    llm_config = load_llm_config()
    return {
        "goal_prompt": GOAL_PROMPT_VERSION,
        "planner_prompt": PLANNER_PROMPT_VERSION,
        "plan_compiler": PLAN_COMPILER_VERSION,
        "capability_registry": get_capability_registry().fingerprint(),
        "source_tree": _source_tree_fingerprint(),
        "golden_suite": hashlib.sha256(
            Path(__file__).with_name("golden_tasks.json").read_bytes()
        ).hexdigest(),
        "model": llm_config.model
        if llm_config is not None
        else os.environ.get("DATA_AGENT_LLM_MODEL")
        or os.environ.get("OPENAI_MODEL")
        or "gpt-4o-mini",
    }


def _source_tree_fingerprint() -> str:
    """Bind an evaluation report to the exact Python implementation that ran it."""

    package_root = Path(__file__).resolve().parents[1]
    digest = hashlib.sha256()
    for source in sorted(package_root.rglob("*.py")):
        digest.update(source.relative_to(package_root).as_posix().encode("utf-8"))
        digest.update(b"\0")
        digest.update(source.read_bytes())
        digest.update(b"\0")
    return digest.hexdigest()


@contextmanager
def _evaluation_environment(
    run_dir: Path,
    *,
    planning_mode: str,
    cache_mode: str,
) -> Iterator[None]:
    overrides = {
        "DATA_AGENT_UNDERSTANDING_CACHE_PATH": str(
            run_dir / "_understanding_cache.json"
        ),
        "DATA_AGENT_UNDERSTANDING_CACHE_ENABLED": "0" if cache_mode == "off" else "1",
    }
    if planning_mode == "deterministic":
        overrides["DATA_AGENT_LLM_ENABLED"] = "0"
    previous = {key: os.environ.get(key) for key in overrides}
    os.environ.update(overrides)
    try:
        yield
    finally:
        for key, value in previous.items():
            if value is None:
                os.environ.pop(key, None)
            else:
                os.environ[key] = value


def _workbook_contains_formula(path: str | Path) -> bool:
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        return any(
            cell.data_type == "f"
            for sheet in workbook.worksheets
            for row in sheet.iter_rows()
            for cell in row
        )
    finally:
        workbook.close()
